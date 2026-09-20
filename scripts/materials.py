#!/usr/bin/env python3
"""Build/check Google Ads preparation workbooks. Requires Python 3.9+ and openpyxl."""
import argparse
import hashlib
import json
import math
import re
import sys
import tempfile
import unicodedata
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlsplit

SPECS = {
    "01": ("关键词规划表", ["序号", "关键词 (Keyword)", "匹配类型建议 (Match Type)", "对应推广页面/产品分类", "排除/挑选原因", "参考链接"]),
    "02": ("目标受众画像分析", ["序号", "受众分类/角色名称", "角色描述", "核心痛点与挑战", "需求与购买动机", "决策关键因素", "针对性营销/行动呼吁建议 (CTA)", "参考链接"]),
    "03": ("谷歌广告描述语", ["序号", "对应页面/产品分类", "英文广告语 (Description)", "字符数自检", "中文翻译", "参考链接"]),
    "04": ("谷歌广告站内链接", ["序号", "对应页面", "Sitelink Text", "Description 1", "Description 2", "Final URL", "参考链接"]),
    "05": ("谷歌广告宣传信息", ["序号", "宣传信息 (Callout Text)", "字符数", "亮点说明", "参考链接"]),
    "06": ("潜在客户表单设置", ["表单字段名称", "设置内容 (英文)", "字符数限制", "实际字符数", "备注", "参考链接"]),
    "07": ("结构化摘要扩展", ["Header 类型", "序号", "属性值 (Value)", "字符数", "参考链接"]),
    "09": ("谷歌广告投放国家推荐", ["推荐层级 (Tier)", "国家名称 (英文)", "国家名称 (中文)", "市场需求/推荐理由", "建议出价策略", "参考链接"]),
}
SEEDS = ["基础词 (Seed Keyword)", "对应产品/页面", "挑选依据", "参考链接"]
REPORT = "08_海外市场与海关宏观报告.html"
REPORT_SECTIONS = ("market-size", "search-trends", "import-markets", "policy-compliance", "competition")
REPORT_FOOTER_LINKS = (
    "https://thinkwithgoogle.com/",
    "https://wits.worldbank.org/",
    "https://www.macmap.org/",
    "https://trends.google.com/trends/",
)
FORBIDDEN = re.compile(r"\b(?:near(?:by)?|me|prices?|pricing|cheap(?:er|est)?|free|old|second[ -]+hand|jobs?|diy|repair(?:ed|ing|s)?)\b", re.I)
CTA = re.compile(r"\b(?:get (?:a |your )?quote|request (?:a |your )?(?:quote|sample|demo|consultation)|contact us|shop now|order now|ask us|talk to|speak to|download (?:the |our |a )?(?:catalog|guide|datasheet)|book (?:a |your )?(?:call|demo|consultation))\b", re.I)
FORM_FIELDS = ["Headline", "Business Name", "Description", "Collected Info", "Privacy Policy URL", "CTA Type", "CTA Description", "Submission Headline", "Submission Description"]
COUNTS = {"03": (2, 3), "05": (1, 2), "06": (1, 3), "07": (2, 3)}
AUDIENCE_DIMENSIONS = {
    "核心受众分类": ("商业模式定位", "目标客户类型", "重点行业"),
    "关键决策者与使用者画像": ("采购与决策角色", "核心痛点与需求", "关键购买因素"),
    "实际应用场景与技术诉求": ("典型应用场景", "技术与功能偏好"),
    "地理与区域市场定位": ("目标区域", "市场定位"),
    "受众总结与营销建议": ("受众总结", "营销建议"),
}
EXTRA_SHEETS = {
    "02": {"audience_analysis": ("五维分析", ["分析维度", "分析项", "分析结论", "依据类型", "证据与推断说明", "参考链接"])},
    "03": {
        "site_hierarchy": ("网站层级", ["页面ID", "页面类型", "层级", "父页面ID", "栏目路径", "Final URL", "文案目标", "证据或缺口说明", "参考链接"]),
        "numeric_claims": ("卖点证据", ["广告语序号", "数据化卖点原文", "指标类型", "指标单位与适用范围", "来源原文摘录", "核验说明", "参考链接"]),
    },
}
PAGE_TYPES = ("品牌与信息", "产品分类", "产品线", "重要页面", "功能页面", "导航分组")
NUMERIC_TYPES = ("产品规格", "产品范围计数", "服务指标", "业务规模", "经验年限")
NUMERIC_VALUE = re.compile(r"\b\d+(?:[.,]\d+)*\b|\b(?:one|two|three|four|five|six|seven|eight|nine|ten|hundred|thousand|dual|triple|single)\b", re.I)
MODEL_ONLY = re.compile(r"[A-Za-z]+[- ]?\d+[A-Za-z0-9]*(?:\s+\d+[A-Za-z]*)?", re.I)
COPY_VIEW_HEADERS = ["栏目与广告语", "中文翻译或说明"]


class ReportStructure(HTMLParser):
    """Inspect report structure, not the truth or relevance of its research."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tags = {}
        self.main_depth = 0
        self.footer_depth = 0
        self.style_depth = 0
        self.css = []
        self.sections = []
        self.section_stack = []
        self.h2_count = 0
        self.footer_links = set()
        self.anchor = None
        self.viewport = False
        self.utf8 = False

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        self.tags[tag] = self.tags.get(tag, 0) + 1
        if tag == "main":
            self.main_depth += 1
        elif tag == "footer":
            self.footer_depth += 1
        elif tag == "style":
            self.style_depth += 1
        elif tag == "section":
            self.sections.append({"id": attrs.get("id"), "in_main": self.main_depth > 0,
                                  "nested": bool(self.section_stack), "h2": 0})
            self.section_stack.append(len(self.sections) - 1)
        elif tag == "h2":
            self.h2_count += 1
            if self.section_stack:
                self.sections[self.section_stack[-1]]["h2"] += 1
        elif tag == "meta":
            self.viewport |= attrs.get("name", "").lower() == "viewport" and bool(attrs.get("content"))
            self.utf8 |= attrs.get("charset", "").lower().replace("-", "") == "utf8"
        elif tag == "a" and self.footer_depth:
            self.anchor = {"href": attrs.get("href", ""), "text": []}

    def handle_endtag(self, tag):
        if tag == "main":
            self.main_depth = max(0, self.main_depth - 1)
        elif tag == "footer":
            self.footer_depth = max(0, self.footer_depth - 1)
        elif tag == "style":
            self.style_depth = max(0, self.style_depth - 1)
        elif tag == "section" and self.section_stack:
            self.section_stack.pop()
        elif tag == "a" and self.anchor is not None:
            if "".join(self.anchor["text"]).strip():
                self.footer_links.add(self.anchor["href"].strip().rstrip("/"))
            self.anchor = None

    def handle_data(self, text):
        if self.style_depth:
            self.css.append(text)
        if self.anchor is not None:
            self.anchor["text"].append(text)


def check_report_html(html):
    parser = ReportStructure()
    parser.feed(html)
    parser.close()
    errors = []
    for tag in ("html", "header", "main", "footer"):
        if parser.tags.get(tag) != 1:
            errors.append(f"08: 需要且仅允许一个{tag}元素")
    if not "".join(parser.css).strip():
        errors.append("08: 缺少内联CSS样式")
    if not parser.viewport:
        errors.append("08: 缺少viewport设置")
    if not parser.utf8:
        errors.append("08: 缺少UTF-8字符集声明")
    if tuple(s["id"] for s in parser.sections) != REPORT_SECTIONS:
        errors.append("08: 正文仅允许五个核心section，依次为 " + ", ".join(REPORT_SECTIONS))
    if parser.h2_count != 5 or any(s["h2"] != 1 for s in parser.sections):
        errors.append("08: 每个核心章节须有且仅有一个h2，共五个；不得增加其他章节标题")
    if any(not s["in_main"] or s["nested"] for s in parser.sections):
        errors.append("08: 五个section须置于main内，不得另建嵌套或外部section")
    for url in REPORT_FOOTER_LINKS:
        if url.rstrip("/") not in parser.footer_links:
            errors.append("08: footer中缺少带可读文字的公开参考链接 " + url)
    return errors


def count(value):
    """Exact for plain English. Double-width text still needs platform review."""
    return sum(2 if unicodedata.east_asian_width(c) in ("W", "F") else 1 for c in str(value or ""))


def keyword_key(value):
    return " ".join(unicodedata.normalize("NFKC", str(value)).split()).casefold()


def seed_sheet_name(index, seed):
    cleaned = re.sub(r"[\[\]:*?/\\\x00-\x1f]", "_", str(seed)).strip().strip("'")
    return (f"词{index:02d}_" + (cleaned or "基础词"))[:31].rstrip("'")


def table_digest(columns, rows):
    payload = json.dumps([columns, rows], ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def is_url(value):
    try:
        p = urlsplit(str(value))
        return p.scheme in ("https", "http") and bool(p.hostname) and not p.username and not re.search(r"\s", str(value))
    except ValueError:
        return False


def filename(step):
    return step + "_" + SPECS[step][0] + ".xlsx"


def schema(steps):
    data = {"sheets": {}}
    for step in steps:
        item = {"filename": filename(step), "columns": SPECS[step][1], "rows": [], "notes": []}
        if step == "01":
            item.update(seed_columns=SEEDS, seeds=[], seed_tables=[])
        for key, (_, headers) in EXTRA_SHEETS.get(step, {}).items():
            item[key] = []
            item[key + "_columns"] = headers
        data["sheets"][step] = item
    return data


def auxiliary_rows(item, step, key, errors):
    """Validate auxiliary tables before business checks or workbook creation."""
    rows = item.get(key)
    label, headers = EXTRA_SHEETS[step][key]
    if not isinstance(rows, list):
        errors.append(f"{step}: 缺少{key}数组（{label}）")
        return []
    valid = []
    for i, row in enumerate(rows, 1):
        if (not isinstance(row, list) or len(row) != len(headers)
                or any(isinstance(v, bool) or not isinstance(v, (str, int)) or not str(v).strip() for v in row)):
            errors.append(f"{step} {label}第{i}行: 需要{len(headers)}个非空文本/整数字段")
            continue
        if any(not is_url(url) for url in str(row[-1]).split("\n")):
            errors.append(f"{step} {label}第{i}行: 参考链接格式错误")
        valid.append(row)
    return valid


def validate_audience(item, errors, warnings, coverage):
    rows = auxiliary_rows(item, "02", "audience_analysis", errors)
    counts, seen, content_seen = {}, set(), set()
    previous_dimension = -1
    for row in rows:
        dimension, topic = row[:2]
        if dimension not in AUDIENCE_DIMENSIONS or topic not in AUDIENCE_DIMENSIONS[dimension]:
            errors.append("02: 未知分析维度或分析项")
            continue
        order = list(AUDIENCE_DIMENSIONS).index(dimension)
        if order < previous_dimension:
            errors.append("02: 五维分析须按指定维度顺序组织")
        previous_dimension = order
        key = (dimension, topic)
        content_key = (dimension, topic, keyword_key(row[2]))
        if content_key in content_seen:
            errors.append("02: 总结、建议或分析内容重复")
        content_seen.add(content_key)
        counts[key] = counts.get(key, 0) + 1
        if key in seen and dimension != "受众总结与营销建议":
            errors.append("02: 分析项重复，请在同一结论内说明相关角色/场景")
        seen.add(key)
        if row[3] not in ("网站事实", "用户提供", "分析推断", "待验证"):
            errors.append("02: 依据类型须区分网站事实、用户提供、分析推断或待验证")
        if row[3] == "待验证":
            warnings.append("02: 待验证分析项：" + topic)
    for dimension, topics in AUDIENCE_DIMENSIONS.items():
        for topic in topics:
            if not counts.get((dimension, topic)):
                errors.append(f"02: 缺少{dimension} / {topic}")
    summaries = counts.get(("受众总结与营销建议", "受众总结"), 0)
    advice = counts.get(("受众总结与营销建议", "营销建议"), 0)
    if not 3 <= summaries <= 4:
        errors.append("02: 受众总结须为3–4行，每行一句")
    if not 2 <= advice <= 3:
        errors.append("02: 营销建议须为2–3条")
    coverage.update(dimensions=len({row[0] for row in rows}), summary_sentences=summaries, marketing_recommendations=advice)


def validate_descriptions(item, rows, gap, errors, warnings, coverage):
    hierarchy = auxiliary_rows(item, "03", "site_hierarchy", errors)
    claims = auxiliary_rows(item, "03", "numeric_claims", errors)
    if not hierarchy:
        errors.append("03: 须先提供网站层级，不能仅交付固定条数广告语")
    pages, paths, targets, urls = {}, set(), {}, set()
    for row in hierarchy:
        page_id, kind, level, parent, path, url, target, reason, _ = row
        if not all(isinstance(v, str) for v in (page_id, kind, parent, path, url, target)):
            errors.append("03: 页面ID、父ID、栏目路径、URL等须为字符串")
            continue
        if page_id in pages or page_id == "ROOT" or path in paths:
            errors.append("03: 页面ID或完整栏目路径重复/无效")
        pages[page_id] = row
        paths.add(path)
        if kind not in PAGE_TYPES or type(level) is not int or level < 0:
            errors.append("03: 页面类型或层级无效")
        if target not in ("生成", "不生成"):
            errors.append("03: 文案目标须为生成或不生成")
        if target == "生成":
            if not is_url(url):
                errors.append("03: 文案目标页面须有真实HTTP(S) URL")
            if url in urls:
                errors.append("03: 同一目标URL不能重复建栏目凑数")
            urls.add(url)
            targets[path] = row
        elif not is_url(url) and not (kind == "导航分组" and url == "无独立页面"):
            errors.append("03: 非页面导航节点须写明无独立页面，其余须有URL")
    for page_id, row in pages.items():
        level, parent = row[2:4]
        if parent == "ROOT":
            if level != 0:
                errors.append("03: ROOT下节点层级须为0")
        elif parent not in pages or parent == page_id:
            errors.append("03: 父页面ID不存在或指向自身")
        elif type(level) is int and type(pages[parent][2]) is int and level != pages[parent][2] + 1:
            errors.append("03: 子层级必须比父节点大1")
    if not targets:
        gap("没有可生成文案的已发现栏目；需继续解析网站")
    ads, grouped = {}, {path: [] for path in targets}
    for row in rows:
        if type(row[0]) is not int or row[0] < 1 or row[0] in ads:
            errors.append("03: 广告语序号必须为唯一正整数")
            continue
        ads[row[0]] = row
        if row[1] not in targets:
            errors.append("03: 广告语栏目不在网站层级的生成目标内：" + str(row[1]))
        else:
            grouped[row[1]].append(row)
    numeric_ids, claim_seen = set(), set()
    for claim in claims:
        ad_id, text, kind = claim[:3]
        if type(ad_id) is not int or ad_id not in ads:
            errors.append("03: 卖点证据引用了不存在的广告语序号")
            continue
        if not isinstance(text, str) or text not in str(ads[ad_id][2]):
            errors.append("03: 数据化卖点须逐字出现在对应英文广告语中")
            continue
        if kind not in NUMERIC_TYPES or not NUMERIC_VALUE.search(text) or MODEL_ONLY.fullmatch(text):
            errors.append("03: 数字卖点须为可量化指标，不能只用型号/标准编号凑数")
            continue
        key = (ad_id, text.casefold())
        if key in claim_seen:
            errors.append("03: 同一广告语的卖点证据重复")
        claim_seen.add(key)
        numeric_ids.add(ad_id)
    per_page = []
    for path, ads_for_page in grouped.items():
        actual = len(ads_for_page)
        numeric = sum(r[0] in numeric_ids for r in ads_for_page)
        ctas = sum(bool(CTA.search(str(r[2]))) for r in ads_for_page)
        per_page.append({"page": path, "target": 5, "actual": actual, "numeric_target": 2, "numeric_actual": numeric, "cta_detected": ctas})
        if actual > 5:
            errors.append(f"03: {path}每个栏目只生成5条，实际{actual}")
        elif actual < 5:
            gap(f"{path}目标5条，实际{actual}，缺口{5-actual}")
        if numeric < 2:
            gap(f"{path}数据化卖点广告目标至少2条，实际{numeric}，缺口{2-numeric}；须补事实证据")
    ctas = sum(bool(CTA.search(str(r[2]))) for r in rows)
    cta_range = (len(rows) // 3, math.ceil(len(rows) / 3))
    coverage.update(target=len(targets) * 5, page_count=len(targets), per_page=per_page,
                    numeric_ads=len(numeric_ids), cta_detected=ctas, cta_target_range=list(cta_range))
    if rows and not cta_range[0] <= ctas <= cta_range[1]:
        gap(f"强CTA目标约1/3（{cta_range[0]}–{cta_range[1]}条），实际检测{ctas}；人工复核表达")


def description_view(item):
    """Readable per-section output; the flat sheet remains the authoritative data."""
    result = []
    for index, page in enumerate((p for p in item["site_hierarchy"] if p[6] == "生成"), 1):
        ads = [r for r in item["rows"] if r[1] == page[4]]
        result.append([f"{index}. {page[4]}", page[5]])
        for number, row in enumerate(ads, 1):
            result.append([f'广告语 {number}: "{row[2]}" (字符数: {row[3]})', row[4]])
        if len(ads) < 5:
            result.append([f"待补充：目标5条，实际{len(ads)}条", page[7]])
    return result


def validate(data, steps, recalculate=False):
    errors, warnings, coverage = [], [], {}
    sheets = data.get("sheets") if isinstance(data, dict) else None
    if not isinstance(sheets, dict):
        return {"errors": ["顶层必须有sheets对象"], "warnings": [], "coverage": {}}
    if set(sheets) - set(SPECS):
        errors.append("存在未知工作簿编号")
    for step in steps:
        item = sheets.get(step)
        if not isinstance(item, dict) or not isinstance(item.get("rows"), list):
            errors.append(step + ": 缺少对象或rows数组")
            continue
        rows, notes = item["rows"], item.get("notes", [])
        if not isinstance(notes, list) or any(not isinstance(n, str) or not n.strip() for n in notes):
            errors.append(step + ": notes必须是非空说明字符串的数组")
            notes = []

        def gap(message):
            (warnings if notes else errors).append(step + ": " + message + ("" if notes else "；须在notes说明原因"))

        coverage[step] = {"actual": len(rows)}
        target = {"01": 30, "02": 1, "04": 20, "05": 10, "09": 1}.get(step)
        if target:
            coverage[step]["target"] = target
            if len(rows) < target or (step in ("03", "04", "05") and len(rows) != target):
                gap(f"目标{target}，实际{len(rows)}，缺口{max(0, target-len(rows))}")
        seen, valid = set(), []
        for i, row in enumerate(rows, 1):
            label = f"{step} 第{i}行"
            if not isinstance(row, list) or len(row) != len(SPECS[step][1]):
                errors.append(label + ": 列数不符")
                continue
            if any(v is not None and (isinstance(v, bool) or not isinstance(v, (str, int, float))) for v in row):
                errors.append(label + ": 单元格仅支持字符串、数字或null")
                continue
            if step in COUNTS:
                src, dest = COUNTS[step]
                actual = count(row[src])
                if recalculate:
                    row[dest] = actual
                elif row[dest] != actual:
                    errors.append(label + ": 字符数与实际不符")
            for j, value in enumerate(row):
                if value is None or str(value).strip() == "":
                    if step == "06" and j == 1:
                        gap(str(row[0]) + "内容缺失，不可直接提交")
                    else:
                        errors.append(label + f": 第{j+1}列为空")
            if any(not is_url(link) for link in str(row[-1] or "").split("\n")):
                errors.append(label + ": 参考链接须为逐行HTTP(S) URL")
            limits = {"03": [(2, 1, 90)], "04": [(2, 1, 25), (3, 1, 35), (4, 1, 35)], "05": [(1, 1, 25)], "07": [(2, 1, 25)]}.get(step, [])
            if step == "06":
                limit = {"Headline": 30, "Description": 200}.get(str(row[0]))
                if isinstance(row[2], int) and row[2] > 0:
                    limit = min(limit, row[2]) if limit else row[2]
                elif row[2] == "平台复核":
                    warnings.append(label + ": 平台限制待核验")
                elif row[2] == "不适用" and row[0] in ("Headline", "Business Name", "Description", "CTA Description", "Submission Headline", "Submission Description"):
                    errors.append(label + ": 文本字段必须填写数值上限或平台复核")
                elif row[2] != "不适用":
                    errors.append(label + ": 限制应为正整数、不适用或平台复核")
                if limit and row[1]:
                    limits = [(1, 1, limit)]
                if row[0] == "Privacy Policy URL" and row[1] and not is_url(row[1]):
                    errors.append(label + ": 隐私URL格式错误")
            for j, low, high in limits:
                value = str(row[j] or "")
                if not low <= count(value) <= high:
                    errors.append(label + f": {SPECS[step][1][j]}计数{count(value)}，要求{low}–{high}")
                if value != value.strip() or any(unicodedata.category(c) in ("Cc", "Cf") for c in value):
                    errors.append(label + ": 文案含首尾空格、换行或不可见控制字符")
                if not value.isascii():
                    warnings.append(label + ": 非ASCII文案需复核平台计数")
            keys = {"01": [1], "02": [1], "03": [2], "04": [2], "05": [1], "06": [0], "07": [0, 2], "09": [1]}[step]
            key = tuple(str(row[j]).strip().casefold() for j in keys)
            if step == "01":
                key = (keyword_key(row[1]),)
            if key in seen:
                errors.append(label + ": 关键内容重复")
            seen.add(key)
            if step == "01":
                if FORBIDDEN.search(str(row[1])):
                    errors.append(label + ": 关键词命中排除词")
                if row[2] not in ("Exact", "Phrase", "Broad"):
                    errors.append(label + ": 匹配类型应为Exact/Phrase/Broad")
            if step == "04" and not is_url(row[5]):
                errors.append(label + ": Final URL格式错误")
            if step == "07" and row[0] not in ("Models", "Service catalog", "Types"):
                errors.append(label + ": 非本流程指定标头")
            if step == "09" and row[0] not in ("Tier 1", "Tier 2", "Tier 3"):
                errors.append(label + ": 层级应为Tier 1/2/3")
            valid.append(row)
        if step == "01":
            seeds = item.get("seeds", [])
            if not isinstance(seeds, list):
                errors.append("01: seeds必须为数组")
                seeds = []
            coverage[step]["seed_actual"] = len(seeds)
            if not seeds:
                gap("至少需要1个用户提供或确认的基础词，实际0")
            seed_seen = set()
            for row in seeds:
                if not isinstance(row, list) or len(row) != 4 or any(not isinstance(v, str) or not v.strip() for v in row):
                    errors.append("01: 基础词须有4个非空字符串字段")
                    continue
                key = keyword_key(row[0])
                if key in seed_seen:
                    errors.append("01: 基础词重复")
                seed_seen.add(key)
                if any(not is_url(x) for x in row[3].split("\n")):
                    errors.append("01: 基础词参考链接格式错误")
            tables = item.get("seed_tables", [])
            if not isinstance(tables, list):
                errors.append("01: seed_tables必须为数组")
                tables = []
            table_seen, source_keywords = set(), set()
            statuses = {"completed": 0, "no_results": 0, "blocked": 0}
            for table in tables:
                if not isinstance(table, dict) or not isinstance(table.get("seed"), str):
                    errors.append("01: 基础词子表缺少seed")
                    continue
                key = keyword_key(table["seed"])
                if key in table_seen or key not in seed_seen:
                    errors.append("01: 基础词子表重复或不在用户基础词清单中")
                table_seen.add(key)
                status = table.get("status")
                headers, detail_rows = table.get("columns"), table.get("rows")
                metadata = table.get("metadata", [])
                if status not in statuses:
                    errors.append("01: 基础词子表status无效")
                    continue
                statuses[status] += 1
                if not isinstance(table.get("context"), str) or not table["context"].strip():
                    errors.append("01: 基础词子表须注明查询条件或受阻原因")
                if not isinstance(metadata, list) or any(not isinstance(r, list) or any(not isinstance(v, str) for v in r) for r in metadata):
                    errors.append("01: CSV元数据须为原始字符串二维数组")
                if status == "blocked":
                    if headers != [] or detail_rows != []:
                        errors.append("01: 受阻子表不得填写指标数据")
                    gap("基础词查询受阻：" + table["seed"])
                    continue
                if not isinstance(table.get("source_file"), str) or not table["source_file"].strip():
                    errors.append("01: 已完成子表缺少来源CSV路径")
                if not isinstance(headers, list) or not headers or any(not isinstance(v, str) for v in headers):
                    errors.append("01: 基础词子表缺少完整原始表头")
                    continue
                col = table.get("keyword_column")
                if type(col) is not int or not 0 <= col < len(headers):
                    errors.append("01: keyword_column须指向真实关键词列（从0开始）")
                    continue
                if not isinstance(detail_rows, list) or any(not isinstance(r, list) or len(r) != len(headers) or any(not isinstance(v, str) for v in r) for r in detail_rows):
                    errors.append("01: 子表指标须为完整原始字符串行，不可转数值或填null")
                    continue
                if (status == "no_results") != (len(detail_rows) == 0):
                    errors.append("01: 子表状态与实际数据行数不一致")
                keys = [keyword_key(r[col]) for r in detail_rows]
                if any(not k for k in keys) or len(keys) != len(set(keys)):
                    errors.append("01: 单个种子子表包含空关键词或重复词")
                source_keywords.update(keys)
            if table_seen != seed_seen:
                errors.append("01: 每个基础词必须恰好对应一个seed_tables子表，受阻也须记录")
            for row in valid:
                if keyword_key(row[1]) not in source_keywords:
                    errors.append("01: 汇总词无法回溯基础词子表：" + str(row[1]))
            coverage[step]["seed_sheet_actual"] = len(tables)
            coverage[step]["seed_statuses"] = statuses
        if step == "02":
            validate_audience(item, errors, warnings, coverage[step])
        if step == "03":
            validate_descriptions(item, valid, gap, errors, warnings, coverage[step])
        if step == "06":
            for name in FORM_FIELDS:
                if name not in {r[0] for r in valid}:
                    gap("缺少字段 " + name)
        if step == "07":
            for header in ("Models", "Service catalog", "Types"):
                actual = sum(r[0] == header for r in valid)
                coverage[step][header] = {"target": 4, "actual": actual}
                if actual != 4:
                    gap(f"{header}目标4，实际{actual}，缺口{max(0, 4-actual)}")
    return {"errors": errors, "warnings": warnings, "coverage": coverage}


def write_sheet(ws, headers, rows):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Website text must never become an Excel formula.
            cell.alignment = Alignment(horizontal="center" if isinstance(cell.value, (int, float)) else "left", vertical="top", wrap_text=True)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="16324F")
                cell.font = Font(bold=True, color="FFFFFF")
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F5FA")
            if isinstance(cell.value, str) and is_url(cell.value):
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
    widths = {}
    for j in range(1, len(headers) + 1):
        widths[j] = min(64, max(12, max(count(line) for row in ws for line in str(row[j-1].value or "").split("\n")) + 2))
        ws.column_dimensions[get_column_letter(j)].width = widths[j]
    for row in ws:
        lines = max(sum(max(1, math.ceil(count(s) / (widths[c.column]-2))) for s in str(c.value or "").split("\n")) for c in row)
        ws.row_dimensions[row[0].row].height = min(360, max(30, lines * 15 + 8))


def build(data, steps, out, replace=False):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    result = validate(data, steps, recalculate=True)
    if result["errors"]:
        return result
    out.mkdir(parents=True, exist_ok=True)
    if not replace and any((out / filename(s)).exists() for s in steps):
        result["errors"].append("输出已存在；确认更新后使用--replace")
        return result
    with tempfile.TemporaryDirectory(prefix=".ads-export-", dir=out) as temp:
        for step in steps:
            item = data["sheets"][step]
            notes = list(item.get("notes", []))
            wb = Workbook()
            wb.active.title = "资料"
            write_sheet(wb.active, SPECS[step][1], item["rows"])
            for key, (title, headers) in EXTRA_SHEETS.get(step, {}).items():
                write_sheet(wb.create_sheet(title), headers, item[key])
            if step == "03":
                ws = wb.create_sheet("栏目文案", 0)
                write_sheet(ws, COPY_VIEW_HEADERS, description_view(item))
                from openpyxl.styles import Font, PatternFill
                for row in ws.iter_rows(min_row=2):
                    if re.match(r"^\d+\. ", str(row[0].value)):
                        for cell in row:
                            cell.fill = PatternFill("solid", fgColor="16324F")
                            cell.font = Font(bold=True, color="FFFFFF")
                wb.active = 0
                notes.append("每个栏目5条是备选描述素材池；单个RSA最多选择4条描述。字符计数包含空格与标点。")
                for page in result["coverage"]["03"]["per_page"]:
                    notes.append(f"{page['page']}：描述目标5、实际{page['actual']}；数据化卖点广告目标至少2、实际{page['numeric_actual']}；强CTA检测{page['cta_detected']}。")
            if step == "02":
                wb.active = wb.sheetnames.index("五维分析")
            if step == "01":
                write_sheet(wb.create_sheet("基础词"), SEEDS, item.get("seeds", []))
                tables = {keyword_key(t["seed"]): t for t in item["seed_tables"]}
                for index, seed_row in enumerate(item["seeds"], 1):
                    table = tables[keyword_key(seed_row[0])]
                    ws = wb.create_sheet(seed_sheet_name(index, seed_row[0]))
                    if table["status"] == "blocked":
                        write_sheet(ws, ["状态", "说明"], [["未取得数据", table["context"]]])
                    else:
                        write_sheet(ws, table["columns"], table["rows"])
                    descriptor = {k: v for k, v in table.items() if k not in ("columns", "rows")}
                    descriptor["format"] = "google-ads-seed-table-v1"
                    descriptor["data_sha256"] = table_digest(table["columns"], table["rows"])
                    ws["A1"].comment = Comment(json.dumps(descriptor, ensure_ascii=False), "Google Ads 来源与查询条件")
                    notes.append(f"基础词：{seed_row[0]} → 子表：{ws.title}；状态：{table['status']}；{table['context']}；来源CSV：{table.get('source_file') or '未取得'}")
            if notes:
                link = next((r[-1].split("\n")[0] for r in item["rows"] if isinstance(r[-1], str)), "https://support.google.com/google-ads/")
                write_sheet(wb.create_sheet("说明"), ["说明", "参考链接"], [[n, link] for n in notes])
            wb.save(Path(temp) / filename(step))
            wb.close()
        for step in steps:
            (Path(temp) / filename(step)).replace(out / filename(step))
    return check(out, steps)


def check(out, steps, require_html=False):
    from openpyxl import load_workbook
    data, file_errors = {"sheets": {}}, []
    for step in steps:
        path = out / filename(step)
        try:
            wb = load_workbook(path, data_only=False)
            values = list(wb["资料"].values)
            if not values or list(values[0]) != SPECS[step][1]:
                file_errors.append(path.name + ": 主表表头错误")
            if any(c.data_type == "f" for ws in wb for row in ws for c in row):
                file_errors.append(path.name + ": 存在公式单元格")
            item = {"rows": [list(r) for r in values[1:]], "notes": []}
            for key, (title, headers) in EXTRA_SHEETS.get(step, {}).items():
                extra = list(wb[title].values)
                if list(extra[0]) != headers:
                    file_errors.append(path.name + ": 附表表头错误 " + title)
                item[key] = [list(r) for r in extra[1:]]
            if step == "03":
                displayed = list(wb["栏目文案"].values)
                if list(displayed[0]) != COPY_VIEW_HEADERS or [list(r) for r in displayed[1:]] != description_view(item):
                    file_errors.append(path.name + ": 栏目文案展示与资料行不一致")
            if "说明" in wb.sheetnames:
                item["notes"] = [r[0] for r in list(wb["说明"].values)[1:] if r[0]]
            if step == "01":
                seeds = list(wb["基础词"].values)
                if not seeds or list(seeds[0]) != SEEDS:
                    file_errors.append(path.name + ": 基础词表头错误")
                item["seeds"] = [list(r) for r in seeds[1:]]
                item["seed_tables"] = []
                for index, seed_row in enumerate(item["seeds"], 1):
                    title = seed_sheet_name(index, seed_row[0])
                    if title not in wb.sheetnames:
                        file_errors.append(path.name + ": 缺少基础词子表 " + title)
                        continue
                    ws = wb[title]
                    if not ws["A1"].comment:
                        file_errors.append(path.name + ": 子表缺少CSV来源与查询元数据 " + title)
                        continue
                    descriptor = json.loads(ws["A1"].comment.text)
                    if descriptor.pop("format", None) != "google-ads-seed-table-v1":
                        file_errors.append(path.name + ": 子表元数据格式错误 " + title)
                    digest = descriptor.pop("data_sha256", None)
                    if descriptor.get("status") == "blocked":
                        headers, detail_rows = [], []
                        if list(ws.values) != [("状态", "说明"), ("未取得数据", descriptor.get("context"))]:
                            file_errors.append(path.name + ": 受阻子表状态被改写 " + title)
                    else:
                        detail = [[str(v) if v is not None else "" for v in r] for r in ws.values]
                        headers, detail_rows = detail[0], detail[1:]
                    if digest != table_digest(headers, detail_rows):
                        file_errors.append(path.name + ": 子表原始指标或列顺序被改写 " + title)
                    if keyword_key(descriptor.get("seed", "")) != keyword_key(seed_row[0]):
                        file_errors.append(path.name + ": 子表基础词与清单不符 " + title)
                    descriptor.update(columns=headers, rows=detail_rows)
                    item["seed_tables"].append(descriptor)
            data["sheets"][step] = item
            wb.close()
        except Exception as exc:
            file_errors.append(path.name + ": 无法读取: " + str(exc))
    result = validate(data, steps)
    result["errors"] = file_errors + result["errors"]
    if require_html:
        try:
            html = (out / REPORT).read_text(encoding="utf-8")
            result["errors"].extend(check_report_html(html))
        except OSError as exc:
            result["errors"].append("08: 无法读取HTML: " + str(exc))
    result["scope"] = "结构、数量、长度、链接、关键词子表、五维受众及逐栏目文案/卖点证据关联；不证明实际查询、全站发现完整性、事实真实性、语法或平台审核通过"
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["schema", "build", "check", "check-html"])
    parser.add_argument("--steps", default=",".join(SPECS), help="默认所有Excel；部分更新可用逗号分隔编号，不含08")
    parser.add_argument("--data", type=Path)
    parser.add_argument("--out", type=Path, default=Path("谷歌广告建组与市场调研资料"))
    parser.add_argument("--replace", action="store_true")
    parser.add_argument("--require-html", action="store_true")
    args = parser.parse_args()
    steps = args.steps.split(",")
    if len(set(steps)) != len(steps) or any(s not in SPECS for s in steps):
        parser.error("--steps必须是不同的01,02,03,04,05,06,07,09编号")
    if args.command == "schema":
        print(json.dumps(schema(steps), ensure_ascii=False, indent=2))
        return 0
    try:
        if args.command == "check-html":
            result = {"errors": check_report_html((args.out / REPORT).read_text(encoding="utf-8")),
                      "warnings": [], "coverage": {},
                      "scope": "08五章结构与页脚链接；研究事实、内容范围及排版须人工核验"}
        elif args.command == "build":
            if not args.data:
                parser.error("build需要--data")
            result = build(json.loads(args.data.read_text(encoding="utf-8")), steps, args.out, args.replace)
        else:
            result = check(args.out, steps, args.require_html)
    except (OSError, ValueError, ImportError) as exc:
        result = {"errors": [str(exc)], "warnings": [], "coverage": {}}
    result["status"] = "invalid" if result["errors"] else "needs_review" if result["warnings"] else "checks_passed"
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return int(bool(result["errors"]))


if __name__ == "__main__":
    sys.exit(main())
