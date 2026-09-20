"""Behavioral tests use synthetic data, never live customer claims or Ads calls."""
import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import materials as m

SOURCE = "https://example.com/products"
GUIDE = "https://support.google.com/google-ads/answer/7684791"


def audience_fixture():
    item = m.schema(["02"])["sheets"]["02"]
    item["rows"] = [[1, "设备OEM工程师", "参与选型与集成", "系统集成成本", "接口可用性", "规格证据", "提供选型资料", SOURCE]]
    for dim, topics in m.AUDIENCE_DIMENSIONS.items():
        for topic in topics:
            for n in range({"受众总结": 3, "营销建议": 2}.get(topic, 1)):
                item["audience_analysis"].append([dim, topic, f"合成测试：{topic}结论{n+1}。", "分析推断", "仅测试结构，无真实客户判断", SOURCE])
    return {"sheets": {"02": item}}


def description_fixture(pages=2):
    item = m.schema(["03"])["sheets"]["03"]
    item["site_hierarchy"] = [["products", "导航分组", 0, "ROOT", "Products", "无独立页面", "不生成", "无独立页面的导航分组", SOURCE]]
    for p in range(pages):
        label = chr(65 + p)
        page_id, path, url = "page-"+label, "Products > Controls " + label, SOURCE + "/" + label.lower()
        item["site_hierarchy"].append([page_id, "产品分类", 1, "products", path, url, "生成", "合成栏目，仅用于测试", url])
        texts = [
            f"Connect 32 digital inputs with {label} controls for cabinet integration.",
            f"Use 2 Ethernet ports on {label} controls for flexible machine networks.",
            f"Choose {label} controls for scalable machine automation.",
            f"Match {label} controls to your cabinet layout and I/O needs.",
            f"Review {label} industrial controls for your equipment project.",
        ]
        for index, text in enumerate(texts):
            number = len(item["rows"]) + 1
            if number % 3 == 0:
                text += " Contact Us."
            item["rows"].append([number, path, text, None, "测试翻译" + str(number), url + "\n" + GUIDE])
            if index < 2:
                phrase = ["32 digital inputs", "2 Ethernet ports"][index]
                item["numeric_claims"].append([number, phrase, "产品规格", "测试范围：该型号接口数", "Synthetic specification: " + phrase, "仅合成测试证据，非真实产品", url])
    return {"sheets": {"03": item}}


class AudienceTests(unittest.TestCase):
    def test_five_dimensions_one_real_role_is_enough(self):
        result = m.validate(audience_fixture(), ["02"], True)
        self.assertEqual(result["errors"], [])
        self.assertEqual(result["warnings"], [])
        self.assertEqual(result["coverage"]["02"]["dimensions"], 5)
        self.assertEqual(result["coverage"]["02"]["summary_sentences"], 3)

    def test_missing_dimension_cannot_be_waived_by_notes(self):
        data = audience_fixture(); item = data["sheets"]["02"]
        item["audience_analysis"] = [r for r in item["audience_analysis"] if r[0] != "地理与区域市场定位"]
        item["notes"] = ["已说明缺口"]
        self.assertTrue(m.validate(data, ["02"], True)["errors"])

    def test_summary_and_advice_counts_and_uniqueness(self):
        for topic in ("受众总结", "营销建议"):
            with self.subTest(topic=topic):
                data = audience_fixture(); rows = data["sheets"]["02"]["audience_analysis"]
                target = next(r for r in rows if r[1] == topic)
                rows.remove(target)
                self.assertTrue(m.validate(data, ["02"], True)["errors"])
        data = audience_fixture(); rows = data["sheets"]["02"]["audience_analysis"]
        rows.append(copy.deepcopy(next(r for r in rows if r[1] == "受众总结")))
        self.assertTrue(m.validate(data, ["02"], True)["errors"])

    def test_pending_evidence_is_visible_as_warning(self):
        data = audience_fixture(); data["sheets"]["02"]["audience_analysis"][0][3] = "待验证"
        result = m.validate(data, ["02"], True)
        self.assertEqual(result["errors"], [])
        self.assertTrue(result["warnings"])


class DescriptionTests(unittest.TestCase):
    def test_five_per_page_and_dynamic_total_accept_short_copy(self):
        for pages in (1, 2, 5):
            with self.subTest(pages=pages):
                data = description_fixture(pages)
                result = m.validate(data, ["03"], True)
                self.assertEqual(result["errors"], [])
                self.assertEqual(result["warnings"], [])
                self.assertEqual(result["coverage"]["03"]["target"], pages * 5)
                self.assertTrue(any(r[3] < 70 for r in data["sheets"]["03"]["rows"]))

    def test_each_page_must_have_its_own_five_ads(self):
        data = description_fixture(); item = data["sheets"]["03"]
        item["rows"][4][1] = item["rows"][5][1]
        result = m.validate(data, ["03"], True)
        self.assertTrue(result["errors"])
        self.assertEqual([p["actual"] for p in result["coverage"]["03"]["per_page"]], [4, 6])

    def test_new_target_with_no_ads_is_not_silently_ignored(self):
        data = description_fixture(); item = data["sheets"]["03"]
        item["site_hierarchy"].append(["about", "品牌与信息", 0, "ROOT", "About", SOURCE + "/about", "生成", "Synthetic page", SOURCE])
        result = m.validate(data, ["03"], True)
        self.assertTrue(result["errors"])
        self.assertEqual(result["coverage"]["03"]["target"], 15)

    def test_ninety_character_boundary_and_count_includes_spaces(self):
        data = description_fixture(); row = data["sheets"]["03"]["rows"][4]
        row[2] = "A " * 44 + "AB"
        self.assertEqual(m.validate(data, ["03"], True)["errors"], [])
        self.assertEqual(row[3], 90)
        row[2] += "C"
        self.assertTrue(m.validate(data, ["03"], True)["errors"])

    def test_numeric_evidence_must_cover_two_distinct_ads_per_page(self):
        data = description_fixture(); item = data["sheets"]["03"]
        item["numeric_claims"] = [r for r in item["numeric_claims"] if r[0] != 2]
        item["numeric_claims"].append([1, "32", "产品规格", "同一条中的另一个片段", "32 inputs", "仅测试", SOURCE])
        result = m.validate(data, ["03"], True)
        self.assertTrue(result["errors"])
        self.assertEqual(result["coverage"]["03"]["per_page"][0]["numeric_actual"], 1)

    def test_numeric_gaps_produce_partial_delivery_only(self):
        data = description_fixture(); item = data["sheets"]["03"]
        item["numeric_claims"] = []; item["notes"] = ["各栏目数字卖点目标2、实际0；待补规格证据。"]
        result = m.validate(data, ["03"], True)
        self.assertEqual(result["errors"], [])
        self.assertTrue(result["warnings"])

    def test_false_evidence_links_and_model_numbers_are_rejected(self):
        for mutation in ("orphan", "absent_phrase", "model", "bad_url"):
            with self.subTest(mutation=mutation):
                data = description_fixture(); item = data["sheets"]["03"]; claim = item["numeric_claims"][0]
                if mutation == "orphan": claim[0] = 999
                elif mutation == "absent_phrase": claim[1] = "500 units"
                elif mutation == "bad_url": claim[-1] = "not a link"
                else:
                    claim[1] = "S7 1200"; item["rows"][0][2] = "Select S7 1200 controls for automation."
                self.assertTrue(m.validate(data, ["03"], True)["errors"])

    def test_cta_target_scales_with_copy_count(self):
        data = description_fixture(); item = data["sheets"]["03"]
        for row in item["rows"]: row[2] = row[2].replace(" Contact Us.", "")
        result = m.validate(data, ["03"], True)
        self.assertTrue(result["errors"])
        self.assertEqual(result["coverage"]["03"]["cta_target_range"], [3, 4])

    def test_invalid_tree_unknown_column_and_duplicate_destination(self):
        for mutation in ("parent", "depth", "destination", "unknown_path", "missing_tree"):
            with self.subTest(mutation=mutation):
                data = description_fixture(); item = data["sheets"]["03"]
                if mutation == "parent": item["site_hierarchy"][1][3] = "missing"
                elif mutation == "depth": item["site_hierarchy"][1][2] = 8
                elif mutation == "destination": item["site_hierarchy"][2][5] = item["site_hierarchy"][1][5]
                elif mutation == "unknown_path": item["rows"][0][1] = "Unmapped"
                else: del item["site_hierarchy"]
                self.assertTrue(m.validate(data, ["03"], True)["errors"])

    def test_old_twenty_row_format_is_not_accepted_as_new_contract(self):
        data = description_fixture(4); item = data["sheets"]["03"]
        del item["site_hierarchy"]; del item["numeric_claims"]
        self.assertTrue(m.validate(data, ["03"], True)["errors"])


class WorkbookTests(unittest.TestCase):
    def test_build_readback_and_presentation_tampering(self):
        data = audience_fixture(); data["sheets"].update(description_fixture()["sheets"])
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp); result = m.build(data, ["02", "03"], out)
            self.assertEqual(result["errors"], [])
            self.assertEqual(result["warnings"], [])
            wb = load_workbook(out / m.filename("02"))
            self.assertEqual(wb.active.title, "五维分析"); wb.close()
            wb = load_workbook(out / m.filename("03"))
            self.assertEqual(wb.active.title, "栏目文案")
            self.assertEqual(wb["网站层级"].max_row, 4)
            self.assertIn('广告语 1: "', wb["栏目文案"]["A3"].value)
            self.assertEqual(wb["资料"]["D2"].value, len(wb["资料"]["C2"].value))
            wb["栏目文案"]["A3"] = "tampered"; wb.save(out / m.filename("03")); wb.close()
            self.assertTrue(m.check(out, ["03"])["errors"])

    def test_build_does_not_overwrite_without_replace(self):
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp); m.build(description_fixture(), ["03"], out)
            before = (out / m.filename("03")).read_bytes()
            self.assertTrue(m.build(description_fixture(), ["03"], out)["errors"])
            self.assertEqual((out / m.filename("03")).read_bytes(), before)

    def test_formula_like_website_text_stays_literal(self):
        data = audience_fixture(); data["sheets"]["02"]["rows"][0][2] = '=HYPERLINK("https://example.com")'
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp); self.assertEqual(m.build(data, ["02"], out)["errors"], [])
            wb = load_workbook(out / m.filename("02")); self.assertEqual(wb["资料"]["C2"].data_type, "s"); wb.close()

    def test_numeric_evidence_readback_cannot_silently_disappear(self):
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp); m.build(description_fixture(), ["03"], out)
            wb = load_workbook(out / m.filename("03")); del wb["卖点证据"]; wb.save(out / m.filename("03")); wb.close()
            self.assertTrue(m.check(out, ["03"])["errors"])


class UnchangedContractTests(unittest.TestCase):
    def test_sitelink_callout_and_snippet_caps(self):
        cases = {
            "04": [1, "Products", "A" * 25, "B" * 35, "C" * 35, SOURCE, SOURCE],
            "05": [1, "A" * 25, None, "test", SOURCE],
            "07": ["Models", 1, "A" * 25, None, SOURCE],
        }
        for step, row in cases.items():
            with self.subTest(step=step):
                data = {"sheets": {step: {"rows": [row], "notes": ["合成数据：已说明数量不足。"]}}}
                self.assertEqual(m.validate(data, [step], True)["errors"], [])
                row[{"04": 2, "05": 1, "07": 2}[step]] += "X"
                self.assertTrue(m.validate(data, [step], True)["errors"])

    def test_lead_form_headline_description_caps_remain(self):
        for field, limit in (("Headline", 30), ("Description", 200)):
            row = [field, "X" * limit, limit, None, "test", SOURCE]
            data = {"sheets": {"06": {"rows": [row], "notes": ["合成数据：仅测试一个字段。"]}}}
            self.assertEqual(m.validate(data, ["06"], True)["errors"], [])
            row[1] += "X"; self.assertTrue(m.validate(data, ["06"], True)["errors"])

    def test_html_still_requires_exactly_five_sections_and_footer_links(self):
        sections = ''.join(f'<section id="{s}"><h2>{s}</h2><p>test</p></section>' for s in m.REPORT_SECTIONS)
        links = ''.join(f'<a href="{u}">source</a>' for u in m.REPORT_FOOTER_LINKS)
        html = f'<html><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width"><style>body{{color:black}}</style></head><body><header>Report</header><main>{sections}</main><footer>{links}</footer></body></html>'
        self.assertEqual(m.check_report_html(html), [])
        self.assertTrue(m.check_report_html(html.replace('</main>', '<section id="extra"><h2>extra</h2></section></main>')))
        self.assertTrue(m.check_report_html(html.replace(m.REPORT_FOOTER_LINKS[0], 'https://example.com/')))


if __name__ == "__main__":
    unittest.main()
